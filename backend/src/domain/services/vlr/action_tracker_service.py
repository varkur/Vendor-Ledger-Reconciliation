"""
Action Tracker domain service.

The Action Tracker tab is the Firmway-style per-reconciliation-request work
queue: every unmatched / residual-difference ledger row across ALL cases in
one reconciliation request, grouped by who needs to act on it (Action Taken
Status: "Pending with Party", "Pending with Company", "No Action Required").

The diagnosis columns (Status, Classification, both-side invoice detail) are
never persisted — they're computed live by reusing the exact same row-building
and labelling logic ReconciliationExportService already uses for the
Reconciliation sheet export, so the Action Tracker screen and the downloaded
workbook can never disagree with the per-case export or with each other.

Only the human-entered workflow fields (action owner, action-taken reference/
remark, request-closed flag) are persisted, in ActionTrackerModel, keyed by a
stable `row_key` derived from the case + the ledger entry id(s) the row
represents.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import and_, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.domain.services.vlr.reconciliation_export_service import (
    ReconciliationExportService,
    derive_action_taken_status,
)
from src.infrastructure.database.models.vlr.action_tracker_model import ActionTrackerModel
from src.infrastructure.database.models.vlr.ledger_entry_model import LedgerEntryModel
from src.infrastructure.database.models.vlr.match_result_model import MatchResultModel
from src.infrastructure.database.models.vlr.reconciliation_case_model import ReconciliationCaseModel
from src.infrastructure.database.models.vlr.reconciliation_request_model import ReconciliationRequestModel
from src.infrastructure.database.models.vlr.vendor_model import VendorModel

# Rows whose Status is one of these are informational (balances / already
# reconciled with nothing outstanding, or self-cancelling knock-off entries)
# and never appear in the Action Tracker — only genuine open items are
# tracked for action. "Reversal Entries" (knock-off/AB pairs) are one-sided
# by construction (each leg nets against another entry on the SAME side,
# not a missing counterpart on the other ledger) — their raw amount is NOT
# an outstanding difference the way an unmatched entry's is, so unlike a
# genuine one-sided unmatched row they must be dropped regardless of amount.
_ALWAYS_NO_ACTION_STATUSES = frozenset({"Reconciled", "Reversal Entries"})

ACTION_TAKEN_STATUSES = ("Pending with Party", "Pending with Company", "No Action Required")

# Column headers for the Action Tracker export, in exact column order —
# matches the reference Firmway export (docs/Action Tracker_MAY-2026-
# ROLEOUT DATA.xlsx) so the downloaded file is a drop-in match.
ACTION_TRACKER_EXPORT_HEADERS = [
    "Party Code", "Party Name", "Company Id", "partyId", "Match Id",
    "Status", "Classification", "Action Taken Status", "Action Owner",
    "Company Invoice Date", "Company Invoice Number", "Company DocType",
    "Company Original DocType", "Company Narration", "Company Amount",
    "Party Invoice Date", "Party Invoice Number", "Party DocType",
    "Party Original DocType", "Party Narration", "Party Amount",
    "Difference", "Remarks", "Request Closed", "Action Taken Reference",
    "Action Taken Remark", "Reco Data & Time", "Posting Date",
    "Clearing date", "Clearing Document Number", "TDS Amount",
]


@dataclass
class ActionTrackerRow:
    """A single Action Tracker detail row (one unmatched/residual entry)."""

    row_key: str
    case_id: str
    vendor_code: str
    vendor_name: str
    company_id: str
    party_id: str
    match_id: str
    status: str
    classification: str
    action_taken_status: str
    company_invoice_date: str = ""
    company_invoice_number: str = ""
    company_doctype: str = ""
    company_original_doctype: str = ""
    company_narration: str = ""
    company_amount: Decimal | str = ""
    party_invoice_date: str = ""
    party_invoice_number: str = ""
    party_doctype: str = ""
    party_original_doctype: str = ""
    party_narration: str = ""
    party_amount: Decimal | str = ""
    difference: Decimal = Decimal("0")
    remarks: str = ""
    reco_datetime: str = ""
    posting_date: str = ""
    clearing_date: str = ""
    clearing_document_number: str = ""
    tds_amount: Decimal | str = ""
    # Human-entered workflow fields (persisted; default = untouched row).
    action_owner: str = "Unassigned"
    action_taken_reference: str = ""
    action_taken_remark: str = ""
    request_closed: bool = False


@dataclass
class ActionTrackerSummaryRow:
    action_taken_status: str
    number_of_records: int
    percentage: float
    amount: Decimal


@dataclass
class ActionTrackerSummary:
    rows: list[ActionTrackerSummaryRow] = field(default_factory=list)
    total_records: int = 0
    total_amount: Decimal = Decimal("0")


def _s(value) -> str:
    """Coerce a possibly-None export row value to a plain string, never None.

    The underlying ledger entry fields (description/narration, invoice
    number, etc.) are nullable, and the row dicts built by
    ReconciliationExportService pass them through as-is (None when absent)
    rather than normalizing to "". The Action Tracker response schema
    declares these as plain `str` fields — Pydantic only applies a field's
    default when the key is OMITTED, not when None is explicitly passed, so
    an unmatched entry with no narration/invoice number raised a 500
    (ValidationError: Input should be a valid string) instead of rendering
    as blank.
    """
    return "" if value is None else str(value)


async def _load_tax_config(case: ReconciliationCaseModel, session: AsyncSession) -> tuple[float, float, float]:
    """Load (tds_max, tds_min, gst_pct) from the case's parent request."""
    if not getattr(case, "request_id", None):
        return 0.0, 0.0, 0.0
    pres = await session.execute(
        select(ReconciliationRequestModel).where(
            ReconciliationRequestModel.id == str(case.request_id)
        )
    )
    parent = pres.scalar_one_or_none()
    if parent is None:
        return 0.0, 0.0, 0.0
    return (
        float(getattr(parent, "tds_percentage", 0) or 0),
        float(getattr(parent, "tds_percentage_min", 0) or 0),
        float(getattr(parent, "gst_percentage", 0) or 0),
    )


class ActionTrackerService:
    """Builds Action Tracker rows for a reconciliation request and manages
    the human-entered workflow fields attached to each row."""

    def __init__(self, session: AsyncSession):
        self._session = session

    @staticmethod
    def _row_key(case_id: str, company_id: str, party_id: str) -> str:
        """
        Stable identity for a row across repeated reads: the case plus
        whichever ledger entry id(s) the row represents. A matched pair row
        carries both ids; a one-sided row carries only the id of the side
        that's present.
        """
        return f"{case_id}:{company_id or '-'}:{party_id or '-'}"

    async def _build_case_rows(
        self,
        case: ReconciliationCaseModel,
        vendor: VendorModel,
    ) -> list[dict]:
        """Build raw export-style row dicts (status/classification/amounts)
        for one case, reusing ReconciliationExportService's own logic."""
        comp_res = await self._session.execute(
            select(LedgerEntryModel).where(
                and_(LedgerEntryModel.case_id == str(case.id), LedgerEntryModel.side == "company")
            )
        )
        company_entries = list(comp_res.scalars().all())
        party_res = await self._session.execute(
            select(LedgerEntryModel).where(
                and_(LedgerEntryModel.case_id == str(case.id), LedgerEntryModel.side == "vendor")
            )
        )
        vendor_entries = list(party_res.scalars().all())

        mr_res = await self._session.execute(
            select(MatchResultModel).where(MatchResultModel.case_id == str(case.id))
        )
        match_results = list(mr_res.scalars().all())

        tds_max, tds_min, gst_pct = await _load_tax_config(case, self._session)
        export_svc = ReconciliationExportService()
        export_svc._tds_percentage_value = tds_max
        export_svc._tds_percentage_min_value = tds_min
        export_svc._gst_percentage_value = gst_pct

        by_id = {str(e.id): e for e in company_entries + vendor_entries}
        return export_svc._build_recon_rows(company_entries, vendor_entries, match_results, by_id)

    async def _get_cases(self, request_id: UUID) -> list[tuple[ReconciliationCaseModel, VendorModel]]:
        res = await self._session.execute(
            select(ReconciliationCaseModel, VendorModel)
            .join(VendorModel, ReconciliationCaseModel.vendor_id == VendorModel.id)
            .where(
                ReconciliationCaseModel.request_id == str(request_id),
                ReconciliationCaseModel.is_deleted == False,  # noqa: E712
            )
        )
        return list(res.all())

    async def _get_workflow_map(self, request_id: UUID) -> dict[str, ActionTrackerModel]:
        """Load all persisted workflow entries for this request, keyed by row_key."""
        res = await self._session.execute(
            select(ActionTrackerModel).where(ActionTrackerModel.request_id == str(request_id))
        )
        return {m.row_key: m for m in res.scalars().all()}

    async def build_rows(self, request_id: UUID) -> list[ActionTrackerRow]:
        """
        Build every Action Tracker row (one per genuinely outstanding
        unmatched/residual entry) across all cases in the request, merged
        with any persisted workflow fields.
        """
        cases = await self._get_cases(request_id)
        workflow_map = await self._get_workflow_map(request_id)

        rows: list[ActionTrackerRow] = []
        for case, vendor in cases:
            raw_rows = await self._build_case_rows(case, vendor)
            for r in raw_rows:
                status = r.get("status") or ""
                if status in _ALWAYS_NO_ACTION_STATUSES:
                    continue
                action_status = derive_action_taken_status(r)

                # The export row dict's own "difference" is 0.0 for any
                # one-sided row by construction (_row() in
                # ReconciliationExportService only computes c_amt + p_amt
                # when BOTH sides are present — see _build_recon_rows).
                # That's correct for the Reconciliation sheet's own
                # Difference column (it represents the matched-pair gap,
                # which is meaningless for a one-sided entry), but Action
                # Tracker's "Difference" is the Firmway-style outstanding
                # amount still unexplained — for a one-sided entry that IS
                # just the entry's own amount (see reference export: Party
                # Amount 12,029 -> Difference 12,029). Recompute it here
                # instead of trusting the row dict's "difference" key so a
                # genuinely outstanding one-sided entry doesn't get
                # reported as a zero-amount, invisible line item.
                is_two_sided = bool(r.get("company_id")) and bool(r.get("party_id"))
                if is_two_sided:
                    diff = Decimal(str(r.get("difference", 0) or 0))
                else:
                    one_sided_amt = r.get("c_amt") if r.get("company_id") else r.get("p_amt")
                    diff = Decimal(str(one_sided_amt or 0))

                # Rows that are genuinely closed out (no residual, no
                # one-sided presence) don't belong on the tracker at all —
                # only Opening/Closing Balance rows and real open items do.
                if action_status == "No Action Required" and abs(diff) < Decimal("0.005") and status not in (
                    "Opening Balance", "Closing Balance",
                ):
                    continue

                row_key = self._row_key(str(case.id), r.get("company_id", ""), r.get("party_id", ""))
                wf = workflow_map.get(row_key)

                rows.append(
                    ActionTrackerRow(
                        row_key=row_key,
                        case_id=str(case.id),
                        vendor_code=vendor.vendor_code or "",
                        vendor_name=vendor.name or "",
                        company_id=_s(r.get("company_id")),
                        party_id=_s(r.get("party_id")),
                        match_id=_s(r.get("matched_id")),
                        status=status,
                        classification=_s(r.get("classification")),
                        action_taken_status=action_status,
                        company_invoice_date=_s(r.get("c_date")),
                        company_invoice_number=_s(r.get("c_invno")),
                        company_doctype=_s(r.get("c_doctype")),
                        company_original_doctype=_s(r.get("c_origtype")),
                        company_narration=_s(r.get("c_narr")),
                        company_amount=_s(r.get("c_amt")),
                        party_invoice_date=_s(r.get("p_date")),
                        party_invoice_number=_s(r.get("p_invno")),
                        party_doctype=_s(r.get("p_doctype")),
                        party_original_doctype=_s(r.get("p_origtype")),
                        party_narration=_s(r.get("p_narr")),
                        party_amount=_s(r.get("p_amt")),
                        difference=diff,
                        remarks=_s(r.get("remark")),
                        reco_datetime=datetime.now(timezone.utc).strftime("%d-%b-%Y %I:%M %p"),
                        posting_date=_s(r.get("c_posting_date")),
                        clearing_date=_s(r.get("c_clearing_date")),
                        clearing_document_number=_s(r.get("c_clearing_doc")),
                        tds_amount=_s(r.get("c_tds_amount")) or "0",
                        action_owner=(wf.action_owner if wf and wf.action_owner else "Unassigned"),
                        action_taken_reference=_s(wf.action_taken_reference if wf else None),
                        action_taken_remark=_s(wf.action_taken_remark if wf else None),
                        request_closed=bool(wf.request_closed) if wf else False,
                    )
                )
        return rows

    async def build_summary(self, request_id: UUID) -> ActionTrackerSummary:
        """Group all Action Tracker rows by action_taken_status for the
        summary screen (counts, percentage, net amount)."""
        rows = await self.build_rows(request_id)
        buckets: dict[str, list] = {s: [0, Decimal("0")] for s in ACTION_TAKEN_STATUSES}
        for r in rows:
            b = buckets.setdefault(r.action_taken_status, [0, Decimal("0")])
            b[0] += 1
            b[1] += r.difference

        total_records = sum(b[0] for b in buckets.values())
        total_amount = sum((b[1] for b in buckets.values()), Decimal("0"))

        summary_rows = []
        for status in ACTION_TAKEN_STATUSES:
            count, amount = buckets.get(status, [0, Decimal("0")])
            if count == 0 and status not in buckets:
                continue
            pct = float(count) / total_records * 100 if total_records else 0.0
            summary_rows.append(
                ActionTrackerSummaryRow(
                    action_taken_status=status,
                    number_of_records=count,
                    percentage=round(pct, 0),
                    amount=amount,
                )
            )
        # Include any bucket not in the canonical ordering (defensive).
        for status, (count, amount) in buckets.items():
            if status in ACTION_TAKEN_STATUSES or count == 0:
                continue
            pct = float(count) / total_records * 100 if total_records else 0.0
            summary_rows.append(
                ActionTrackerSummaryRow(
                    action_taken_status=status,
                    number_of_records=count,
                    percentage=round(pct, 0),
                    amount=amount,
                )
            )

        return ActionTrackerSummary(
            rows=summary_rows,
            total_records=total_records,
            total_amount=total_amount,
        )

    async def update_item(
        self,
        request_id: UUID,
        case_id: UUID,
        row_key: str,
        *,
        action_owner: str | None = None,
        action_taken_reference: str | None = None,
        action_taken_remark: str | None = None,
        request_closed: bool | None = None,
        updated_by: str = "system",
    ) -> ActionTrackerModel:
        """Create or update the persisted workflow fields for one row."""
        res = await self._session.execute(
            select(ActionTrackerModel).where(
                and_(
                    ActionTrackerModel.case_id == str(case_id),
                    ActionTrackerModel.row_key == row_key,
                )
            )
        )
        entry = res.scalar_one_or_none()
        if entry is None:
            entry = ActionTrackerModel(
                request_id=str(request_id),
                case_id=str(case_id),
                row_key=row_key,
                created_by=updated_by,
                modified_by=updated_by,
            )
            self._session.add(entry)

        if action_owner is not None:
            entry.action_owner = action_owner
        if action_taken_reference is not None:
            entry.action_taken_reference = action_taken_reference
        if action_taken_remark is not None:
            entry.action_taken_remark = action_taken_remark
        if request_closed is not None:
            entry.request_closed = request_closed
        entry.modified_by = updated_by

        await self._session.flush()
        return entry

    async def build_export(self, request_id: UUID, request_title: str = "") -> bytes:
        """
        Build the Action Tracker detail grid as an .xlsx workbook with the
        exact 31-column layout of the reference Firmway export.
        """
        rows = await self.build_rows(request_id)

        wb = Workbook()
        ws = wb.active
        ws.title = "Action Tracker"

        header_fill = PatternFill("solid", fgColor="7CB342")
        header_font = Font(bold=True, color="FFFFFF")
        for col_idx, header in enumerate(ACTION_TRACKER_EXPORT_HEADERS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, r in enumerate(rows, start=2):
            values = [
                r.vendor_code,
                r.vendor_name,
                r.company_id,
                r.party_id,
                r.match_id,
                r.status,
                r.classification,
                r.action_taken_status,
                r.action_owner,
                r.company_invoice_date,
                r.company_invoice_number,
                r.company_doctype,
                r.company_original_doctype,
                r.company_narration,
                r.company_amount,
                r.party_invoice_date,
                r.party_invoice_number,
                r.party_doctype,
                r.party_original_doctype,
                r.party_narration,
                r.party_amount,
                float(r.difference),
                r.remarks,
                "Yes" if r.request_closed else "No",
                r.action_taken_reference,
                r.action_taken_remark,
                r.reco_datetime,
                r.posting_date,
                r.clearing_date,
                r.clearing_document_number,
                r.tds_amount,
            ]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Reasonable column widths so the file is readable without manual
        # resizing (mirrors the reference export's autosize behaviour).
        for col_idx, header in enumerate(ACTION_TRACKER_EXPORT_HEADERS, start=1):
            width = max(12, min(32, len(header) + 4))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    async def bulk_update_items(
        self,
        request_id: UUID,
        items: list[tuple[UUID, str]],
        *,
        action_owner: str | None = None,
        action_taken_reference: str | None = None,
        action_taken_remark: str | None = None,
        request_closed: bool | None = None,
        updated_by: str = "system",
    ) -> int:
        """Apply the same workflow update to multiple (case_id, row_key) rows."""
        count = 0
        for case_id, row_key in items:
            await self.update_item(
                request_id,
                case_id,
                row_key,
                action_owner=action_owner,
                action_taken_reference=action_taken_reference,
                action_taken_remark=action_taken_remark,
                request_closed=request_closed,
                updated_by=updated_by,
            )
            count += 1
        return count
